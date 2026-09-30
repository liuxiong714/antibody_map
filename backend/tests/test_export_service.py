"""services/analysis/export.py unit tests (pure openpyxl, no DB)."""
from __future__ import annotations

import io

from openpyxl import load_workbook

from app.services.analysis.export import build_excel_export


def _wb(blob: bytes):
    return load_workbook(io.BytesIO(blob))


def test_full_pipeline_returns_nonempty_bytes():
    summary = {"total": 5000, "studies": 25, "rate": 18.5}
    trend = [{"year": 2019, "rate": 12.3, "n": 1200}]
    region = [{"province": "Sichuan", "rate": 22.1, "n": 800}]
    age = [{"age_group": "0-4", "rate": 2.1, "n": 300}]
    approved = [{
        "literature_title": "T", "disease": "measles", "sample_size": 50,
        "value": 0.85, "literature_authors": "A", "literature_journal": "J",
        "literature_year": 2022, "literature_doi": "10.x/y",
    }]
    blob = build_excel_export(trend, region, age, summary, approved)
    assert isinstance(blob, bytes) and len(blob) > 1024


def test_empty_data_still_returns_workbook():
    blob = build_excel_export([], [], [], {}, [])
    assert isinstance(blob, bytes)


def test_non_dict_inputs_safe():
    blob = build_excel_export([], [], [], ["a", "b"], [])
    assert isinstance(blob, bytes)


def test_trend_rows_not_dict():
    blob = build_excel_export([(2019, 12.3)], [], [], {}, [])
    assert isinstance(blob, bytes)


def test_region_adds_sheet():
    region = [{"province": "Sichuan", "rate": 22.1}]
    blob = build_excel_export([], region, [], {}, [])
    wb = _wb(blob)
    assert len(wb.sheetnames) >= 3  # summary + region + methods


def test_approved_without_trend_region_age():
    approved = [{"literature_title": "T", "value": 0.7}]
    blob = build_excel_export([], [], [], {}, approved)
    wb = _wb(blob)
    assert len(wb.sheetnames) >= 3  # summary + dp_detail + methods


def test_methods_appendix_present():
    blob = build_excel_export([], [], [], {}, [])
    wb = _wb(blob)
    assert len(wb.sheetnames) >= 2  # summary + methods at minimum


def test_empty_trend_region_age_no_sheets():
    blob = build_excel_export([], [], [], {}, [])
    wb = _wb(blob)
    assert len(wb.sheetnames) == 2  # summary + methods only
